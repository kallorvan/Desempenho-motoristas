#!/usr/bin/env python3
"""Gera saida/Indicadores_Produtividade_Motoristas_MM_AAAA.xlsx a partir de dados/AAAA-MM.json.

Uso: python3 scripts/2_planilha.py AAAA-MM

Replica o modelo (docs/procedimento.md, "Modelo da planilha") e recalcula as
fórmulas com o LibreOffice, conferindo que nenhuma célula ficou com erro.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / 'scripts'))
from comum import CODIGO_GRUPO, GRUPOS, SEM_GRUPO  # noqa: E402

CAB = ['Cod. Motorista(Rodopar)', 'Motorista', 'Grupo de Maior KM', 'Descrição Grupo',
       'Total Frete Cliente (R$)', 'Total Bonus Operacional (R$)', 'Valor Economia (R$)',
       'Premio por Economia (R$)', '% Bonus por Media', 'Bonus por Media (R$)',
       'Total Premiacao (R$)', 'Total Bonificações (R$)', 'Qtd Viagens', 'Qtd Viagem+Vira',
       'KM Rodado', 'Media (km/l)', 'Posicao no Grupo', '% Bonificacao Sobre Total Frete Cliente',
       'Frete Medio p/ Viagem (R$)', 'Bonificacao Media p/ Viagem (R$)', 'Frete por KM (R$/km)',
       'Bonificacao por KM (R$/km)', 'Part. % no Frete Total', 'Ranking Frete', 'Ranking Bonificacao']

MOEDA = '#,##0.00'
FMT = {'E': MOEDA, 'F': MOEDA, 'G': MOEDA, 'H': MOEDA, 'I': '0%', 'J': MOEDA, 'K': MOEDA, 'L': MOEDA,
       'M': '0', 'N': '0', 'O': '#,##0', 'P': '0.0000', 'Q': '0', 'R': '0.00%', 'S': MOEDA,
       'T': MOEDA, 'U': '0.0000', 'V': '0.0000', 'W': '0.00%', 'X': '0', 'Y': '0'}
AZUL, PRETO, BRANCO = '0000FF', '000000', 'FFFFFF'
VERDE, NAVY, TOT = '00B050', '1F3864', 'D9E1F2'
FINO = Side(style='thin', color='BFBFBF')
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def gcod(r):
    if 'gcod' in r:
        return r['gcod']
    return CODIGO_GRUPO[r['grupo']]  # bases de jul/ago não guardam o código


def montar(mes):
    rows = json.loads((RAIZ / 'dados' / f'{mes}.json').read_text(encoding='utf8'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Base Motorista'

    # --- aba de classificação: todos os grupos conhecidos (+ sem grupo, se houver)
    tab = sorted(GRUPOS.items())
    if any(gcod(r) == SEM_GRUPO[0] for r in rows):
        tab = [SEM_GRUPO] + tab
    wc = wb.create_sheet('Base Classificação Motorista')
    for i, (c, n) in enumerate(tab, 1):
        wc.cell(i, 1, c).font = Font(name='Arial', size=10)
        wc.cell(i, 2, n).font = Font(name='Arial', size=10)
    wc.column_dimensions['B'].width = 38
    faixa = f"'Base Classificação Motorista'!$A$1:$B${len(tab)}"

    # --- cabeçalho
    for j, h in enumerate(CAB, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name='Arial', size=10, bold=True, color=BRANCO)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[1].height = 42

    n = len(rows)
    ult, tot = n + 1, n + 2
    for i, d in enumerate(rows, 2):
        cod = d['cod'] if re.fullmatch(r'\d+', d['cod']) else str(d['cod']).replace('#', '')
        vals = {
            'A': cod, 'B': d['nome'], 'C': gcod(d),
            'D': f'=VLOOKUP(C{i},{faixa},2,FALSE())',
            'E': d['frete'], 'F': d['bonusOp'], 'G': d['econ'], 'H': d['premio'],
            'I': d['pctMedia'], 'J': d['bMedia'],
            'K': f'=H{i}+J{i}', 'L': f'=F{i}+K{i}',
            'M': d['viag'], 'N': d['vira'], 'O': d['km'], 'P': d['media'], 'Q': d['pos'],
            'R': f'=IFERROR(L{i}/E{i},"Manobrista")', 'S': f'=IFERROR(E{i}/M{i},"")',
            'T': f'=IFERROR(L{i}/N{i},"")', 'U': f'=IFERROR(E{i}/O{i},"")',
            'V': f'=IFERROR(L{i}/O{i},"")', 'W': f'=IFERROR(E{i}/$E${tot},"")',
            'X': f'=IF(E{i}=0,"",RANK(E{i},$E$2:$E${ult}))',
            'Y': f'=IF(L{i}=0,"",RANK(L{i},$L$2:$L${ult}))',
        }
        for col, v in vals.items():
            c = ws[f'{col}{i}']
            c.value = v
            verde = col in 'LR'
            calc = col in 'DKSTUVWXY'
            c.font = Font(name='Arial', size=10, color=BRANCO if verde else (PRETO if calc else AZUL))
            if verde:
                c.fill = PatternFill('solid', fgColor=VERDE)
            if col in FMT:
                c.number_format = FMT[col]
            c.border = BORDA

    # --- linha de total
    t = {'A': 'TOTAL / MEDIA', 'B': f'{n} Motoristas'}
    for col in 'EFGHJKLMNO':
        t[col] = f'=SUM({col}2:{col}{ult})'
    t['P'] = f'=AVERAGEIFS(P2:P{ult},P2:P{ult},">0")'
    t['R'] = f'=IFERROR(L{tot}/E{tot},"Manobrista")'
    t['S'] = f'=IFERROR(E{tot}/M{tot},"")'
    t['T'] = f'=IFERROR(L{tot}/N{tot},"")'
    t['U'] = f'=IFERROR(E{tot}/O{tot},"")'
    t['V'] = f'=IFERROR(L{tot}/O{tot},"")'
    t['W'] = f'=IFERROR(E{tot}/$E${tot},"")'
    for j in range(1, 26):
        col = ws.cell(1, j).column_letter
        c = ws[f'{col}{tot}']
        c.value = t.get(col)
        c.font = Font(name='Arial', size=11, bold=True, color=BRANCO if col in 'LR' else PRETO)
        c.fill = PatternFill('solid', fgColor=VERDE if col in 'LR' else TOT)
        if col in FMT:
            c.number_format = FMT[col]
        c.border = BORDA

    larg = {'A': 12, 'B': 40, 'C': 10, 'D': 36}
    for j in range(1, 26):
        col = ws.cell(1, j).column_letter
        ws.column_dimensions[col].width = larg.get(col, 15)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:Y{ult}'
    return wb, rows


def recalcular(xlsx):
    """Abre e salva no LibreOffice para gravar os resultados das fórmulas no arquivo."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                        '--outdir', tmp, str(xlsx)], check=True, capture_output=True)
        shutil.copy(Path(tmp) / xlsx.name, xlsx)


def conferir(xlsx, rows):
    wb = load_workbook(xlsx, data_only=True)
    ws = wb['Base Motorista']
    erros = [c.coordinate for r in ws.iter_rows() for c in r
             if isinstance(c.value, str) and c.value.startswith('#')]
    vazias = [f'{col}{i}' for i in range(2, len(rows) + 2) for col in 'DKL' if ws[f'{col}{i}'].value is None]
    tot = len(rows) + 2
    return erros, vazias, {k: ws[f'{c}{tot}'].value for k, c in
                           (('frete', 'E'), ('premTot', 'K'), ('bonif', 'L'), ('media', 'P'), ('pct', 'R'))}


def main():
    if len(sys.argv) != 2 or not re.fullmatch(r'\d{4}-\d{2}', sys.argv[1]):
        sys.exit(__doc__)
    mes = sys.argv[1]
    aaaa, mm = mes.split('-')
    wb, rows = montar(mes)
    (RAIZ / 'saida').mkdir(exist_ok=True)
    xlsx = RAIZ / 'saida' / f'Indicadores_Produtividade_Motoristas_{mm}_{aaaa}.xlsx'
    wb.save(xlsx)
    recalcular(xlsx)
    erros, vazias, t = conferir(xlsx, rows)
    print(f'Gravado: {xlsx.relative_to(RAIZ)} — {len(rows)} motoristas')
    print(f'Total: frete {t["frete"]:,.2f} · premiação {t["premTot"]:,.2f} · bonificações {t["bonif"]:,.2f} · '
          f'média {t["media"]:.4f} · % bonif/frete {t["pct"]:.4%}')
    if erros or vazias:
        sys.exit(f'ERRO: células com erro {erros[:20]} · fórmulas sem resultado {vazias[:20]}')
    print('Fórmulas recalculadas, nenhuma célula com erro.')


if __name__ == '__main__':
    main()
