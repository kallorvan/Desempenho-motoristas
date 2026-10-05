#!/usr/bin/env python3
"""Confronta o export xlsx do Rodopar ("Comissão Toliman.rpt") com o PDF do mês.

Uso: python3 scripts/conferir_xlsx.py entrada/<export>.xlsx entrada/<arquivo>.pdf AAAA-MM

O xlsx traz os mesmos lançamentos e o RESUMO DO BÔNUS, com valores sem
arredondamento, mas não traz o RESUMO PRÊMIO POR MÉDIA. Serve para separar
divergência real da fonte de efeito de leitura/arredondamento do PDF.
Precisa de saida/conferencia_AAAA-MM.json (gerado pelo 1_extrair.py).
"""
import importlib.util
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location('extrair', RAIZ / 'scripts' / '1_extrair.py')
ex = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ex)


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, str):
        v = v.strip()
        return float(v.replace('.', '').replace(',', '.')) if v else 0.0
    return float(v)


def ler_xlsx(path):
    """{cod: {nome, linhas, resumo}} — colunas do export: I frete, L VG, M lona(qtd),
    N vira, O carregamento, P lona, Q viagem (sem arredondar)."""
    ws = load_workbook(path, data_only=True).active
    drv, cur = {}, None
    for r in ws.iter_rows(values_only=True):
        a = r[0]
        if isinstance(a, str) and a.startswith('MOTORISTA:'):
            m = re.match(r'MOTORISTA:\s*(.*?)\s+-\s+COD\.:\s*(\S+)', a)
            cur = drv.setdefault(m.group(2), {'nome': m.group(1), 'l': [], 'res': {}})
        elif isinstance(a, str) and re.match(r'[MOE]: ', a):
            cur['l'].append(dict(doc=a, frete=num(r[8]), vg=int(num(r[11])), vira=num(r[13]),
                                 carreg=num(r[14]), lona=num(r[15]), viagem=num(r[16])))
        elif a == 'RESUMO DO MOTORISTA':
            # pares rótulo/valor; rótulo repetido = 1º quantidade, 2º valor
            vistos = defaultdict(int)
            for i in range(1, len(r) - 1):
                if isinstance(r[i], str) and r[i].endswith(':'):
                    k = r[i] + ('#v' if vistos[r[i]] else '')
                    vistos[r[i]] += 1
                    cur['res'][k] = num(r[i + 1])
    return drv


def main():
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    xlsx, pdf, mes = sys.argv[1:]
    X = ler_xlsx(xlsx)
    conf = json.loads((RAIZ / 'saida' / f'conferencia_{mes}.json').read_text(encoding='utf8'))
    C = {r['cod']: r for r in conf['rows']}
    txt = ex.texto_pdf(pdf)
    ms = list(ex.RE_MOTORISTA.finditer(txt))
    PL = defaultdict(list)
    for i, m in enumerate(ms):
        fim = ms[i + 1].start() if i + 1 < len(ms) else len(txt)
        PL[m.group(2)].extend(ex.lancamentos(txt[m.end():fim].split('\n'))[0])

    print(f'== xlsx × PDF — {mes}')
    print(f'Motoristas: xlsx {len(X)} · PDF {len(C)} · só no xlsx {sorted(set(X) - set(C))} · '
          f'só no PDF {sorted(set(C) - set(X))}')
    print(f'Lançamentos: xlsx {sum(len(v["l"]) for v in X.values())} · PDF {sum(len(v) for v in PL.values())}')

    mapa = (('frete', 'TOTAL DE FRETE CLIENTE:', 'frete'), ('_vViag', 'TOTAL DE VIAGENS:#v', 'viagem'),
            ('_vVira', 'TOTAL DE VIRAS:#v', 'vira'), ('_vCarreg', 'TOTAL CARREGAMENTO:#v', 'carreg'),
            ('_vLona', 'TOTAL DE LONAS:#v', 'lona'), ('viag', 'TOTAL DE VIAGENS:', None),
            ('vira', 'TOTAL VIAGEM+VIRA:', None), ('bonusOp', 'SUBTOTAL BÔNUS:', None))
    res, lin, fonte = [], [], []
    for cod in sorted(set(X) & set(C)):
        x, p = X[cod], C[cod]
        for kp, kx, kl in mapa:
            vx = x['res'].get(kx, 0.0)
            if abs(round(vx + 1e-9, 2) - p[kp]) > 0.005:
                res.append(f'{cod} {x["nome"]}: {kx} PDF {p[kp]:,.2f} × xlsx {vx:,.4f}')
            if kl and abs(sum(l[kl] for l in x['l']) - vx) > 0.005:
                fonte.append(f'{cod} {x["nome"]}: {kx.rstrip(":#v")} resumo {vx:,.2f} × linhas '
                             f'{sum(l[kl] for l in x["l"]):,.2f} (dif. {vx - sum(l[kl] for l in x["l"]):+,.2f})')
        P = PL[cod]
        if len(P) != len(x['l']):
            lin.append(f'{cod}: {len(P)} linhas no PDF × {len(x["l"])} no xlsx')
        for a, b in zip(P, x['l']):
            for k in ('frete', 'vira', 'carreg', 'lona', 'viagem'):
                tol = 0.0051 if k == 'viagem' else 0.005  # o PDF arredonda o bônus de cada linha
                if abs(a[k] - b[k]) > tol:
                    lin.append(f'{cod} {b["doc"]}: {k} PDF {a[k]:,.2f} × xlsx {b[k]:,.4f}')

    for t, lst in (('Resumos PDF × xlsx diferentes', res), ('Linhas PDF × xlsx diferentes', lin),
                   ('Resumo × soma das linhas no xlsx (divergência da fonte, sem arredondamento)', fonte)):
        print(f'\n-- {t}: {len(lst)}')
        for e in lst:
            print('   ' + e)


if __name__ == '__main__':
    main()
