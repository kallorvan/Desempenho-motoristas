#!/usr/bin/env python3
"""Compara duas versões do PDF "Planilha de comissão" do mesmo mês.

Uso: python3 scripts/comparar_versoes.py <pdf_pago> <pdf_novo> AAAA-MM

<pdf_pago> é a versão usada no pagamento (a do fechamento); <pdf_novo> é a reemitida.
Gera saida/Comparativo_Versoes_PDF_MM_AAAA.xlsx e imprime o resumo:
- motoristas com valor diferente (frete, bônus, premiação, bonificações, viagens) e
  o acerto (novo − pago: negativo = pago a mais, positivo = pago a menos);
- lançamentos alterados, excluídos e incluídos, campo a campo;
- mudanças só de Nº RV (administrativas, não mexem em valor);
- cruzamento com a apuração de pagamento indevido (regras.apuracao) do mês.

Compara os PDFs como vieram: não aplica dados/ajustes (vale o resumo de cada versão).
"""
import importlib.util
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / 'scripts'))
import regras  # noqa: E402

_spec = importlib.util.spec_from_file_location('extrair', RAIZ / 'scripts' / '1_extrair.py')
extrair = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(extrair)
CAMPOS = extrair.CAMPOS_LANC

# valores do motorista comparados (rótulo, chave)
VALORES = [('Frete cliente', 'frete'), ('Bônus operacional', 'bonusOp'), ('Total premiação', 'premTot'),
           ('Total bonificações', 'bonif'), ('Viagens', 'viag'), ('Viagem+vira', 'vira'), ('KM', 'km'),
           ('Média', 'media')]
# campos do lançamento comparados lado a lado
CAMPOS_LADO = [('Frete', 'frete'), ('Nº RV', 'rv'), ('VG', 'vg'), ('Lona (qtd)', 'lonaq'), ('Vira', 'vira'),
               ('Carreg.', 'carreg'), ('Lona R$', 'lona'), ('Viagem', 'viagem'), ('Trecho', 'od')]
CHAVE = ('doc', 'cte', 'container', 'data')


def bonus(x):
    return round(x['vira'] + x['carreg'] + x['lona'] + x['viagem'], 2)


def ler(pdf):
    rows, movs, rel, _ = extrair.processar(pdf)
    falhas = rel['falhas']
    if falhas:
        sys.exit(f'{pdf.name}: {len(falhas)} linha(s)/bloco(s) não lidos — corrigir o parser antes:\n' +
                 '\n'.join(falhas[:10]))
    return {r['cod']: r for r in rows}, movs


def comparar_lanc(a_ls, b_ls):
    """Pareia os lançamentos de um motorista. Devolve (alterados, excluídos, incluídos, só RV)."""
    tup = lambda x: tuple(x[k] for k in CAMPOS)  # noqa: E731
    ca, cb = Counter(map(tup, a_ls)), Counter(map(tup, b_ls))
    sa = [dict(zip(CAMPOS, t)) for t in (ca - cb).elements()]
    sb = [dict(zip(CAMPOS, t)) for t in (cb - ca).elements()]
    # ordem do PDF
    ordem = {tup(x): i for i, x in enumerate(a_ls)}
    sa.sort(key=lambda x: ordem.get(tup(x), 0))
    alt, so_rv = [], []
    for a in list(sa):
        par = next((b for b in sb if all(b[k] == a[k] for k in CHAVE)), None)
        if par is None:  # mesma linha com outra chave: tenta doc + CT-e
            par = next((b for b in sb if b['doc'] == a['doc'] and b['cte'] == a['cte']), None)
        if par is None:
            continue
        sa.remove(a)
        sb.remove(par)
        dif = [k for k in CAMPOS if a[k] != par[k]]
        (so_rv if dif == ['rv'] else alt).append((a, par, dif))
    return alt, sa, sb, so_rv


def descrever(a, b, dif):
    """Leitura curta da alteração de um lançamento."""
    t = []
    if a['viagem'] and not b['viagem']:
        t.append(f'viagem de R$ {a["viagem"]:,.2f} retirada')
    elif a['viagem'] != b['viagem']:
        pa = a['viagem'] / a['frete'] * 100 if a['frete'] else 0
        pb = b['viagem'] / b['frete'] * 100 if b['frete'] else 0
        t.append(f'viagem R$ {a["viagem"]:,.2f} → {b["viagem"]:,.2f}' +
                 (f' (taxa {pa:.2f}% → {pb:.2f}%)' if a['frete'] and a['frete'] == b['frete'] else ''))
    if a['carreg'] != b['carreg']:
        t.append(f'carregamento R$ {a["carreg"]:,.2f} → {b["carreg"]:,.2f}')
    if a['frete'] != b['frete']:
        t.append(f'frete R$ {a["frete"]:,.2f} → {b["frete"]:,.2f}')
    if a['vg'] != b['vg']:
        t.append(f'VG {a["vg"]} → {b["vg"]}')
    if a['lona'] != b['lona'] or a['lonaq'] != b['lonaq']:
        t.append(f'lona {a["lonaq"]} / R$ {a["lona"]:,.2f} → {b["lonaq"]} / R$ {b["lona"]:,.2f}')
    if a['vira'] != b['vira']:
        t.append(f'vira R$ {a["vira"]:,.2f} → {b["vira"]:,.2f}')
    if 'rv' in dif:
        t.append(f'Nº RV {a["rv"] or "—"} → {b["rv"] or "—"}')
    for k in ('od', 'cliente', 'container', 'chegada'):
        if k in dif:
            t.append(f'{k} {a[k]} → {b[k]}')
    return '; '.join(t)


def comparar(pdf_a, pdf_b, mes):
    ra, ma = ler(pdf_a)
    rb, mb = ler(pdf_b)
    cods = sorted(set(ra) | set(rb), key=lambda c: (ra.get(c) or rb.get(c))['nome'])
    mot, lanc, rv = [], [], []
    for c in cods:
        a, b = ra.get(c), rb.get(c)
        nome = (a or b)['nome']
        alt, exc, inc, so_rv = comparar_lanc(ma.get(c, []), mb.get(c, []))
        for x, y, dif in alt:
            lanc.append(dict(cod=c, nome=nome, sit='Alterado', a=x, b=y, dif=dif, txt=descrever(x, y, dif),
                             db=round(bonus(y) - bonus(x), 2)))
        for x in exc:
            lanc.append(dict(cod=c, nome=nome, sit='Excluído na versão nova', a=x, b=None, dif=CAMPOS,
                             txt=f'lançamento retirado (bônus R$ {bonus(x):,.2f})', db=-bonus(x)))
        for y in inc:
            lanc.append(dict(cod=c, nome=nome, sit='Incluído na versão nova', a=None, b=y, dif=CAMPOS,
                             txt=f'lançamento novo (bônus R$ {bonus(y):,.2f})', db=bonus(y)))
        for x, y, _ in so_rv:
            rv.append(dict(cod=c, nome=nome, a=x, b=y))
        vals = {k: ((a or {}).get(k, 0) or 0, (b or {}).get(k, 0) or 0) for _, k in VALORES}
        mudou = [k for k, (u, v) in vals.items() if abs(u - v) > 0.0001]
        if mudou or alt or exc or inc or a is None or b is None:
            mot.append(dict(cod=c, nome=nome, vals=vals, mudou=mudou, nAlt=len(alt), nExc=len(exc), nInc=len(inc),
                            sit='Só na versão paga' if b is None else 'Só na versão nova' if a is None else ''))
    tot = {k: (round(sum(r[k] for r in ra.values()), 2), round(sum(r[k] for r in rb.values()), 2)) for _, k in VALORES
           if k not in ('media',)}
    tot['motoristas'] = (len(ra), len(rb))
    tot['lancamentos'] = (sum(map(len, ma.values())), sum(map(len, mb.values())))
    nomes = {c: r['nome'] for c, r in (rb | ra).items()}
    return dict(mot=mot, lanc=lanc, rv=rv, tot=tot, nomes=nomes, apur=cruzar_apuracao(mes, lanc))


def cruzar_apuracao(mes, lanc):
    """A apuração de pagamento indevido (feita na versão paga) × o que a versão nova mudou.
    Agrupa por motorista + documento; as diferenças entram pelo documento ou pelo CT-e
    (na ordem cancelada a nova pode ter retirado a outra ordem do mesmo CT-e)."""
    if not regras.carregar(mes):
        return None
    ap = regras.apuracao(mes, regras.verificar(mes))
    grupos = {}
    for it in ap['itens']:
        grupos.setdefault((it['cod'], it['doc']), []).append(it)
    linhas, cobertos = [], set()
    for (cod, doc), its in grupos.items():
        ctes = {it['cte'] for it in its}
        hits = [x for x in lanc if x['cod'] == cod and ((x['a'] or x['b'])['doc'] == doc or (x['a'] or x['b'])['cte'] in ctes)]
        cobertos.update(id(h) for h in hits)
        rec = round(sum(it['recuperar'] for it in its), 2)
        recup = round(-sum(h['db'] for h in hits), 2)
        sobra_viagem = [h for h in hits if h['b'] and h['b']['viagem'] and h['b']['vg'] == 0]
        if not hits:
            st = 'Mantido — não corrigido na versão nova'
        elif sobra_viagem:
            st = 'Parcial — ainda há viagem paga'
        else:
            st = 'Corrigido na versão nova'
        nota = '; '.join(dict.fromkeys(h['txt'] for h in hits))
        if hits and abs(rec - recup) > 0.004:
            nota += (f' — a nova recupera R$ {recup - rec:,.2f} a mais que a apuração' if recup > rec else
                     f' — a nova recupera R$ {rec - recup:,.2f} a menos que a apuração')
        linhas.append(dict(cod=cod, doc=doc, motivo=its[0]['motivo'], data=its[0]['data'], od=its[0]['od'],
                           cte=' / '.join(it['cte'] for it in its),
                           container=' / '.join(it['container'] for it in its if it['container']),
                           n=len(its), pago=round(sum(it['pago'] for it in its), 2),
                           devido=round(sum(it['devido'] for it in its), 2), recuperar=rec, recup=recup,
                           pend=round(rec - recup, 2), st=st, nota=nota))
    fora = [x for x in lanc if id(x) not in cobertos]
    return dict(itens=linhas, fora=fora, total=ap['total'])


# ---------------------------------------------------------------- planilha
AZUL = '1F3864'
F_CAB = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FILL_CAB = PatternFill('solid', fgColor=AZUL)
F = Font(name='Arial', size=10)
FB = Font(name='Arial', size=10, bold=True)
F_TIT = Font(name='Arial', size=13, bold=True, color=AZUL)
F_NOTA = Font(name='Arial', size=9, italic=True, color='595959')
FILL_DIF = PatternFill('solid', fgColor='FFF2CC')
FILL_A = PatternFill('solid', fgColor='F2F2F2')
FILL_TOT = PatternFill('solid', fgColor='D9E1F2')
F_NEG = Font(name='Arial', size=10, bold=True, color='C00000')
F_POS = Font(name='Arial', size=10, bold=True, color='007A3D')
BORDA = Border(bottom=Side(style='thin', color='BFBFBF'))
MOEDA = '#,##0.00;[Red]-#,##0.00'
MOEDA_S = '+#,##0.00;[Red]-#,##0.00;0.00'


def cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(linha, i, t)
        c.font, c.fill = F_CAB, FILL_CAB
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[linha].height = 32
    ws.freeze_panes = ws.cell(linha + 1, 3)
    ws.auto_filter.ref = f'A{linha}:{get_column_letter(len(titulos))}{linha}'
    for i, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, r, c, v, fmt=None, font=F, fill=None):
    cell = ws.cell(r, c, v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    cell.border = BORDA
    return cell


def acerto_font(v):
    return F_NEG if v < -0.004 else F_POS if v > 0.004 else F


def planilha(res, pdf_a, pdf_b, mes, dest):
    wb = Workbook()
    aaaa, mm = mes.split('-')

    # --- Resumo
    ws = wb.active
    ws.title = 'Resumo'
    ws['A1'] = f'Comparativo de versões do PDF — fechamento {mm}/{aaaa}'
    ws['A1'].font = F_TIT
    ws['A2'] = f'Versão paga: {pdf_a.name}'
    ws['A3'] = f'Versão nova: {pdf_b.name}'
    ws['A4'] = (f'Gerado em {date.today():%d/%m/%Y}. Acerto = versão nova − versão paga: negativo = pago a mais '
                '(a descontar), positivo = pago a menos (a pagar).')
    for k in ('A2', 'A3'):
        ws[k].font = F
    ws['A4'].font = F_NOTA
    cabecalho(ws, 6, ['Indicador', 'Versão paga', 'Versão nova', 'Diferença'], [34, 18, 18, 18])
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    linhas = [('Motoristas', 'motoristas', '0'), ('Lançamentos', 'lancamentos', '0')] + \
             [(rot, k, '0' if k in ('viag', 'vira', 'km') else MOEDA) for rot, k in VALORES if k in res['tot']]
    r = 7
    for rot, k, fmt in linhas:
        u, v = res['tot'][k]
        put(ws, r, 1, rot, font=FB)
        put(ws, r, 2, u, fmt)
        put(ws, r, 3, v, fmt)
        put(ws, r, 4, round(v - u, 2), MOEDA_S if fmt == MOEDA else '+0;-0;0',
            font=acerto_font(v - u) if k in ('bonusOp', 'premTot', 'bonif') else FB)
        r += 1
    r += 1
    nl = Counter(x['sit'] for x in res['lanc'])
    acerto = round(sum(m['vals']['bonif'][1] - m['vals']['bonif'][0] for m in res['mot']), 2)
    menos = round(sum(m['vals']['bonif'][1] - m['vals']['bonif'][0] for m in res['mot']
                      if m['vals']['bonif'][1] < m['vals']['bonif'][0]), 2)
    mais = round(acerto - menos, 2)
    for rot, v, fmt in [
            ('Motoristas com algum valor diferente', len(res['mot']), '0'),
            ('… com bonificação menor na versão nova (pago a mais)',
             sum(1 for m in res['mot'] if m['vals']['bonif'][1] < m['vals']['bonif'][0] - 0.004), '0'),
            ('… com bonificação maior na versão nova (pago a menos)',
             sum(1 for m in res['mot'] if m['vals']['bonif'][1] > m['vals']['bonif'][0] + 0.004), '0'),
            ('Pago a mais (a descontar)', menos, MOEDA_S),
            ('Pago a menos (a pagar)', mais, MOEDA_S),
            ('Acerto líquido (nova − paga)', acerto, MOEDA_S),
            ('Lançamentos alterados', nl['Alterado'], '0'),
            ('Lançamentos excluídos na versão nova', nl['Excluído na versão nova'], '0'),
            ('Lançamentos incluídos na versão nova', nl['Incluído na versão nova'], '0'),
            ('Lançamentos com só o Nº RV alterado (sem valor)', len(res['rv']), '0')]:
        put(ws, r, 1, rot, font=FB)
        put(ws, r, 4, v, fmt, font=acerto_font(v) if fmt == MOEDA_S else FB)
        r += 1
    r += 1
    ws.cell(r, 1, 'Abas: Por motorista (paga × nova e acerto) · Lançamentos (campo a campo, diferenças em amarelo) · '
                  'Nº RV (só documento) · x Apuração (pagamento indevido apurado × o que a versão nova corrigiu).'
            ).font = F_NOTA

    # --- Por motorista
    ws = wb.create_sheet('Por motorista')
    tit = ['Cód.', 'Motorista']
    for rot, _ in VALORES[:6]:
        tit += [f'{rot}\npaga', f'{rot}\nnova', f'{rot}\ndif.']
    tit += ['Acerto\n(nova − paga)', 'Lanç.\nalterados', 'Lanç.\nexcluídos', 'Lanç.\nincluídos', 'O que mudou']
    cabecalho(ws, 1, tit, [7, 34] + [12, 12, 11] * 6 + [14, 10, 10, 10, 90])
    r = 2
    for m in sorted(res['mot'], key=lambda m: m['vals']['bonif'][1] - m['vals']['bonif'][0]):
        put(ws, r, 1, m['cod'])
        put(ws, r, 2, m['nome'] + (f' ({m["sit"]})' if m['sit'] else ''))
        c = 3
        for _, k in VALORES[:6]:
            u, v = m['vals'][k]
            fmt = '0' if k in ('viag', 'vira') else MOEDA
            fill = FILL_DIF if k in m['mudou'] else None
            put(ws, r, c, u, fmt, fill=fill)
            put(ws, r, c + 1, v, fmt, fill=fill)
            put(ws, r, c + 2, round(v - u, 2), MOEDA_S if fmt == MOEDA else '+0;-0;0', fill=fill)
            c += 3
        ac = round(m['vals']['bonif'][1] - m['vals']['bonif'][0], 2)
        put(ws, r, c, ac, MOEDA_S, font=acerto_font(ac))
        put(ws, r, c + 1, m['nAlt'], '0')
        put(ws, r, c + 2, m['nExc'], '0')
        put(ws, r, c + 3, m['nInc'], '0')
        txt = [x['txt'] for x in res['lanc'] if x['cod'] == m['cod']]
        if not txt:
            txt = ['só o resumo mudou: ' + ', '.join(
                f'{rot.lower()} {m["vals"][k][0]:,.2f} → {m["vals"][k][1]:,.2f}' for rot, k in VALORES if k in m['mudou'])]
        put(ws, r, c + 4, ' | '.join(dict.fromkeys(txt)))
        r += 1
    put(ws, r, 2, 'TOTAL', font=FB, fill=FILL_TOT)
    for c in range(3, 3 + 18 + 1):
        col = get_column_letter(c)
        put(ws, r, c, f'=SUM({col}2:{col}{r - 1})', ws.cell(2, c).number_format, font=FB, fill=FILL_TOT)

    # --- Lançamentos
    ws = wb.create_sheet('Lançamentos')
    tit = ['Cód.', 'Motorista', 'Situação', 'Documento', 'Data', 'Container', 'CT-e', 'Cliente']
    for rot, _ in CAMPOS_LADO:
        tit += [f'{rot}\npaga', f'{rot}\nnova']
    tit += ['Bônus\npaga', 'Bônus\nnova', 'Acerto\n(nova − paga)', 'O que mudou']
    larg = [7, 30, 14, 15, 10, 16, 13, 26] + [11, 11] * (len(CAMPOS_LADO) - 1) + [30, 30] + [11, 11, 13, 70]
    cabecalho(ws, 1, tit, larg)
    r = 2
    for x in res['lanc']:
        a, b = x['a'], x['b']
        ref = a or b
        put(ws, r, 1, x['cod'])
        put(ws, r, 2, x['nome'])
        put(ws, r, 3, x['sit'], font=FB if x['sit'] != 'Alterado' else F)
        for c, k in enumerate(('doc', 'data', 'container', 'cte', 'cliente'), 4):
            put(ws, r, c, ref[k])
        c = 9
        for _, k in CAMPOS_LADO:
            fill = FILL_DIF if (k in x['dif'] or x['sit'] != 'Alterado') else None
            fmt = MOEDA if k in ('frete', 'vira', 'carreg', 'lona', 'viagem') else None
            put(ws, r, c, a[k] if a else None, fmt, fill=fill or FILL_A)
            put(ws, r, c + 1, b[k] if b else None, fmt, fill=fill)
            c += 2
        put(ws, r, c, bonus(a) if a else 0, MOEDA, fill=FILL_A)
        put(ws, r, c + 1, bonus(b) if b else 0, MOEDA)
        put(ws, r, c + 2, x['db'], MOEDA_S, font=acerto_font(x['db']))
        put(ws, r, c + 3, x['txt'])
        r += 1
    put(ws, r, 2, 'TOTAL', font=FB, fill=FILL_TOT)
    c0 = 9 + 2 * len(CAMPOS_LADO)
    for c in range(c0, c0 + 3):
        col = get_column_letter(c)
        put(ws, r, c, f'=SUM({col}2:{col}{r - 1})', MOEDA_S, font=FB, fill=FILL_TOT)
    ws.cell(r + 2, 1, 'Colunas "paga" em cinza; células em amarelo = campo diferente entre as versões. '
                      'Bônus = viagem + vira + carregamento + lona da linha. A soma das linhas pode diferir em '
                      'centavos do acerto por motorista, que vem do resumo de cada versão.').font = F_NOTA

    # --- Nº RV
    ws = wb.create_sheet('Nº RV')
    cabecalho(ws, 1, ['Cód.', 'Motorista', 'Documento', 'Data', 'Container', 'CT-e', 'Trecho', 'Nº RV\npaga',
                      'Nº RV\nnova'], [7, 30, 15, 10, 16, 13, 36, 12, 12])
    for r, x in enumerate(res['rv'], 2):
        a, b = x['a'], x['b']
        for c, v in enumerate((x['cod'], x['nome'], a['doc'], a['data'], a['container'], a['cte'], a['od'],
                               a['rv'] or '—', b['rv'] or '—'), 1):
            put(ws, r, c, v, fill=FILL_DIF if c >= 8 else None)
    ws.cell(len(res['rv']) + 3, 1, 'Mudança só no número do documento (Nº RV): nenhum valor muda.').font = F_NOTA

    # --- x Apuração
    ap = res['apur']
    if ap:
        ws = wb.create_sheet('x Apuração')
        ws['A1'] = 'Apuração de pagamento indevido (feita sobre a versão paga) × o que a versão nova corrigiu'
        ws['A1'].font = FB
        ws['A2'] = ('Recuperado pela nova = bônus que a versão nova tirou desses lançamentos. Diferença = a recuperar '
                    '(apuração, carregamento a R$ 25,00) − recuperado pela nova: positivo = a nova ainda deixa valor '
                    'pago a mais; negativo = a nova tirou mais que a apuração.')
        ws['A2'].font = F_NOTA
        tit = ['Cód.', 'Motorista', 'Motivo (apuração)', 'Documento', 'Data', 'Trecho', 'CT-e', 'Container',
               'Lanç.', 'Pago', 'Devido\n(apuração)', 'A recuperar\n(apuração)', 'Recuperado\npela nova',
               'Diferença', 'Situação na versão nova', 'O que a versão nova mudou']
        cabecalho(ws, 4, tit, [7, 30, 26, 15, 10, 30, 26, 32, 7, 11, 11, 13, 13, 12, 36, 90])
        r = 5
        for it in ap['itens']:
            vals = (it['cod'], res['nomes'].get(it['cod'], ''), it['motivo'], it['doc'], it['data'], it['od'],
                    it['cte'], it['container'], it['n'], it['pago'], it['devido'], it['recuperar'], it['recup'],
                    it['pend'], it['st'], it['nota'])
            for c, v in enumerate(vals, 1):
                put(ws, r, c, v, '0' if c == 9 else MOEDA if 10 <= c <= 13 else MOEDA_S if c == 14 else None,
                    font=FB if c == 15 else F_NEG if c == 14 and v > 0.004 else F,
                    fill=FILL_DIF if c == 15 and not it['st'].startswith('Corrigido') else None)
            r += 1
        put(ws, r, 2, 'TOTAL', font=FB, fill=FILL_TOT)
        for c in range(9, 15):
            col = get_column_letter(c)
            put(ws, r, c, f'=SUM({col}5:{col}{r - 1})', '0' if c == 9 else MOEDA if c < 14 else MOEDA_S,
                font=FB, fill=FILL_TOT)
        r += 3
        ws.cell(r, 1, 'Correções da versão nova fora da apuração de pagamento indevido').font = FB
        r += 1
        for c, t in enumerate(['Cód.', 'Motorista', 'Situação', 'Documento', 'Data', 'Trecho', 'CT-e', 'Container',
                               '', 'Bônus paga', 'Bônus nova', '', '', 'Acerto', '', 'O que mudou'], 1):
            cc = ws.cell(r, c, t)
            cc.font, cc.fill = F_CAB, FILL_CAB
        r += 1
        for x in ap['fora']:
            ref = x['a'] or x['b']
            vals = {1: x['cod'], 2: x['nome'], 3: x['sit'], 4: ref['doc'], 5: ref['data'], 6: ref['od'], 7: ref['cte'],
                    8: ref['container'], 10: bonus(x['a']) if x['a'] else 0, 11: bonus(x['b']) if x['b'] else 0,
                    14: x['db'], 16: x['txt']}
            for c, v in vals.items():
                put(ws, r, c, v, MOEDA if c in (10, 11) else MOEDA_S if c == 14 else None,
                    font=acerto_font(v) if c == 14 else F)
            r += 1

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    dest.parent.mkdir(exist_ok=True)
    wb.save(dest)


def main():
    if len(sys.argv) != 4 or not re.fullmatch(r'\d{4}-\d{2}', sys.argv[3]):
        sys.exit(__doc__)
    pdf_a, pdf_b, mes = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
    res = comparar(pdf_a, pdf_b, mes)
    aaaa, mm = mes.split('-')
    dest = RAIZ / 'saida' / f'Comparativo_Versoes_PDF_{mm}_{aaaa}.xlsx'
    planilha(res, pdf_a, pdf_b, mes, dest)
    _spec2 = importlib.util.spec_from_file_location('planilha', RAIZ / 'scripts' / '2_planilha.py')
    pl = importlib.util.module_from_spec(_spec2)
    _spec2.loader.exec_module(pl)
    pl.recalcular(dest)  # grava o resultado dos totais (SUM) no arquivo

    print(f'== Versões do PDF {mes}: {pdf_a.name} (paga) × {pdf_b.name} (nova)')
    for k, (u, v) in res['tot'].items():
        if u != v:
            print(f'   {k}: {u:,.2f} → {v:,.2f} ({v - u:+,.2f})')
    print(f'\n-- Motoristas com valor diferente: {len(res["mot"])}')
    for m in sorted(res['mot'], key=lambda m: m['vals']['bonif'][1] - m['vals']['bonif'][0]):
        ac = m['vals']['bonif'][1] - m['vals']['bonif'][0]
        print(f'   {m["cod"]} {m["nome"]}: acerto {ac:+,.2f} · ' +
              ', '.join(f'{k} {m["vals"][k][0]:,.2f}→{m["vals"][k][1]:,.2f}' for k in m['mudou']))
    print(f'\n-- Lançamentos: {len(res["lanc"])}')
    for x in res['lanc']:
        ref = x['a'] or x['b']
        print(f'   {x["cod"]} {x["sit"]} {ref["doc"]} {ref["data"]} CT-e {ref["cte"]} {ref["od"]}: {x["txt"]} '
              f'(acerto {x["db"]:+,.2f})')
    print(f'\n-- Só Nº RV alterado: {len(res["rv"])} lançamentos')
    if res['apur']:
        ap = res['apur']
        st = Counter(i['st'].split(' ')[0] for i in ap['itens'])
        print(f'\n-- Apuração de pagamento indevido × versão nova: ' + ', '.join(f'{k}: {v}' for k, v in st.items()))
        for i in ap['itens']:
            print(f'   {i["cod"]} {i["doc"]} {i["od"]}: a recuperar {i["recuperar"]:,.2f} · recuperado pela nova '
                  f'{i["recup"]:,.2f} · diferença {i["pend"]:+,.2f} — {i["st"]}')
        print(f'   Total: a recuperar {sum(i["recuperar"] for i in ap["itens"]):,.2f} · recuperado pela nova '
              f'{sum(i["recup"] for i in ap["itens"]):,.2f}')
        print(f'   Correções da nova fora da apuração: {len(ap["fora"])} lançamentos, acerto '
              f'{sum(x["db"] for x in ap["fora"]):+,.2f}')
    print(f'\nGravado: {dest.relative_to(RAIZ)}')


if __name__ == '__main__':
    main()
